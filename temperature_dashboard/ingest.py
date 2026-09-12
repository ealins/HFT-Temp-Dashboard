from __future__ import annotations

import csv
import io
import re
import sqlite3
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile

import pandas as pd
from openpyxl import load_workbook

from .db import (
    DATA_DIR,
    apply_mapping,
    canonical_building,
    canonical_room,
    create_indexes,
    init_db,
    insert_rows,
    summary,
    metadata_set,
)

TEMP_SHEET = re.compile(r"^(?P<building>[^_]+)_(?P<room>.+)_(?P<sensor>HOBO_.+)$", re.I)
DEFAULT_CO2_MAPPING = DATA_DIR / "co2_location_mapping.csv"


def ingest_temperature_excel(file_obj, filename: str = "temperature.xlsx") -> dict:
    with NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(file_obj.read())
        temp_path = Path(tmp.name)
    
    con = init_db("Temperature", replace=True, with_indexes=False)
    con.execute("PRAGMA synchronous=OFF")
    try:
        con.execute("PRAGMA journal_mode=MEMORY")
    except sqlite3.OperationalError:
        pass

    try:
        wb = load_workbook(temp_path, read_only=True, data_only=True)
        try:
            def rows():
                for sheet in wb.sheetnames:
                    m = TEMP_SHEET.match(sheet)
                    if m:
                        building = canonical_building(m.group("building"))
                        room = canonical_room(m.group("room"))
                        sensor = m.group("sensor")
                    else:
                        building, room, sensor = "Unassigned", "Unassigned", sheet
                    for row in wb[sheet].iter_rows(values_only=True):
                        if len(row) < 3:
                            continue
                        ts, raw = row[1], row[2]
                        if ts is None or raw is None:
                            continue
                        try:
                            val = float(raw)
                            if abs(val) > 200:
                                val /= 1000.0
                            stamp = pd.Timestamp(ts).isoformat()
                        except Exception:
                            continue
                        yield (stamp, sensor, building, room, val, None, None, "Temperature Excel")

            total = insert_rows(con, rows())
            create_indexes(con)
        finally:
            wb.close()
        
        con.close()
        out = summary("Temperature")
        out["ingested"] = total
        out["filename"] = filename
        return out
    finally:
        try:
            con.close()
        except Exception:
            pass
        temp_path.unlink(missing_ok=True)


def _sensor_from_name(name: str) -> str:
    stem = Path(name).stem
    stem = re.sub(r"^(CO2sensors_|LoRa_CO2sensors_)", "", stem, flags=re.I)
    return stem or "Unknown"


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _iter_co2_csv_bytes(data: bytes, name: str):
    text = _decode(data)
    stream = io.StringIO(text)
    header = None
    for _ in range(16):
        line = stream.readline()
        if not line:
            break
        low = line.lower()
        if "server time" in low and "co2 concentration" in low:
            header = line
            break
    if header is None:
        return
    reader = csv.DictReader(io.StringIO(header + stream.read()), delimiter=";")
    sensor = _sensor_from_name(name)
    ts_re = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}")

    def f(v):
        try:
            return float(str(v).replace(",", "."))
        except Exception:
            return None

    for row in reader:
        ts_raw = row.get("Server time") or row.get("# Server time") or row.get("Sensor time") or row.get("TTN UTC time")
        co2_raw = row.get("CO2 concentration") or row.get("CO2") or row.get("CO2 concentration (ppm)")
        if ts_raw is None or co2_raw is None:
            continue
        ts_text = str(ts_raw).strip()
        m = ts_re.match(ts_text)
        if not m:
            continue
        stamp = m.group(0).replace(" ", "T")
        try:
            val = float(str(co2_raw).replace(",", "."))
            if not (0 <= val <= 100000):
                continue
        except Exception:
            continue
        temp = f(row.get("Temperature"))
        hum = f(row.get("Humidity"))
        yield (stamp, sensor, "Unassigned", "Unassigned", val, temp, hum, "CO2 upload")


def _iter_co2_csv_file(filepath: Path, name: str):
    """Iterate rows from a CO2 CSV file on disk."""
    text = filepath.read_text(encoding="utf-8-sig", errors="replace")
    stream = io.StringIO(text)
    header = None
    for _ in range(16):
        line = stream.readline()
        if not line:
            break
        low = line.lower()
        if "server time" in low and "co2 concentration" in low:
            header = line
            break
    if header is None:
        return
    reader = csv.DictReader(io.StringIO(header + stream.read()), delimiter=";")
    sensor = _sensor_from_name(name)
    ts_re = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}")

    def f(v):
        try:
            return float(str(v).replace(",", "."))
        except Exception:
            return None

    for row in reader:
        ts_raw = row.get("Server time") or row.get("# Server time") or row.get("Sensor time") or row.get("TTN UTC time")
        co2_raw = row.get("CO2 concentration") or row.get("CO2") or row.get("CO2 concentration (ppm)")
        if ts_raw is None or co2_raw is None:
            continue
        ts_text = str(ts_raw).strip()
        m = ts_re.match(ts_text)
        if not m:
            continue
        stamp = m.group(0).replace(" ", "T")
        try:
            val = float(str(co2_raw).replace(",", "."))
            if not (0 <= val <= 100000):
                continue
        except Exception:
            continue
        temp = f(row.get("Temperature"))
        hum = f(row.get("Humidity"))
        yield (stamp, sensor, "Unassigned", "Unassigned", val, temp, hum, "CO2 upload")


def _apply_packaged_co2_mapping() -> dict:
    if not DEFAULT_CO2_MAPPING.exists():
        return {"mapping_rows": 0, "mapped_readings": 0}
    m = pd.read_csv(DEFAULT_CO2_MAPPING, dtype=str)
    matched = apply_mapping("CO₂", m)
    metadata_set("CO₂", "packaged_hft_mapping", "202510")
    return {"mapping_rows": len(m), "mapped_readings": matched}


def ingest_co2_upload(file_obj, filename: str) -> dict:
    content = file_obj.read()
    con = init_db("CO₂", replace=True, with_indexes=False)
    con.execute("PRAGMA synchronous=OFF")
    try:
        con.execute("PRAGMA journal_mode=MEMORY")
    except sqlite3.OperationalError:
        pass
    try:
        def all_rows():
            low = filename.lower()
            if low.endswith(".zip"):
                with zipfile.ZipFile(io.BytesIO(content)) as z:
                    for name in z.namelist():
                        if name.lower().endswith(".csv"):
                            yield from _iter_co2_csv_bytes(z.read(name), name)
            elif low.endswith(".csv"):
                yield from _iter_co2_csv_bytes(content, filename)
            elif low.endswith((".xlsx", ".xls")):
                xl = pd.ExcelFile(io.BytesIO(content))
                for sheet in xl.sheet_names:
                    df = pd.read_excel(xl, sheet_name=sheet)
                    cols = {str(c).lower(): c for c in df.columns}

                    def pick(words):
                        for k, c in cols.items():
                            if all(w in k for w in words):
                                return c
                        return None

                    tcol = pick(["time"]) or pick(["date"])
                    ccol = pick(["co2"])
                    scol = pick(["sensor"])
                    tempcol = pick(["temp"])
                    hcol = pick(["humid"])
                    if tcol is None or ccol is None:
                        continue
                    sensor_default = sheet
                    for _, r in df.iterrows():
                        try:
                            ts = pd.to_datetime(r[tcol], errors="raise")
                            val = float(r[ccol])
                        except Exception:
                            continue
                        sensor = str(r[scol]) if scol is not None and pd.notna(r[scol]) else sensor_default
                        temp = float(r[tempcol]) if tempcol is not None and pd.notna(r[tempcol]) else None
                        hum = float(r[hcol]) if hcol is not None and pd.notna(r[hcol]) else None
                        yield (pd.Timestamp(ts).isoformat(), sensor, "Unassigned", "Unassigned", val, temp, hum, "CO2 upload")
            else:
                raise ValueError("CO₂ upload must be ZIP, CSV or Excel.")

        total = insert_rows(con, all_rows())
        create_indexes(con)
        con.close()
        mapping_result = _apply_packaged_co2_mapping()
        out = summary("CO₂")
        out["ingested"] = total
        out["filename"] = filename
        out.update(mapping_result)
        return out
    finally:
        try:
            con.close()
        except Exception:
            pass
def ingest_co2_from_path(path: str | Path) -> dict:
    """
    Ingest all CO2 CSV files from a directory path (e.g., mapped network drive W:).
    This avoids browser upload and works with large datasets.
    """
    root = Path(path)
    if not root.exists():
        raise FileNotFoundError(f"Path does not exist: {path}")
    
    # Use os.listdir for better compatibility with WebDAV mounts on Windows
    import os
    try:
        files = os.listdir(str(root))
    except Exception as e:
        raise ValueError(f"Could not list directory {path}: {e}")
    
    csv_files = [root / f for f in files if f.lower().endswith(".csv")]
    if not csv_files:
        raise ValueError(f"No CSV files found in {path}")
    
    con = init_db("CO₂", replace=True, with_indexes=False)
    con.execute("PRAGMA synchronous=OFF")
    try:
        con.execute("PRAGMA journal_mode=MEMORY")
    except sqlite3.OperationalError:
        pass
    
    try:
        def all_rows():
            for csv_file in csv_files:
                yield from _iter_co2_csv_file(csv_file, csv_file.name)
        
        total = insert_rows(con, all_rows())
        create_indexes(con)
        con.close()
        mapping_result = _apply_packaged_co2_mapping()
        out = summary("CO₂")
        out["ingested"] = total
        out["filename"] = f"Directory: {path}"
        out.update(mapping_result)
        return out
    finally:
        try:
            con.close()
        except Exception:
            pass
